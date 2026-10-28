"""Generate a styled Excel (.xlsx) version of the Section 10 Precedents Database.

Exports:
- data/section10_cases_structured_database.xlsx
- /Users/ananyasridhar/Desktop/Section10_Cases_Structured_Database.xlsx
"""

import csv
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_excel(csv_path="data/section10_cases_structured_database.csv", output_path="data/section10_cases_structured_database.xlsx"):
    if not os.path.exists(csv_path):
        raise FileNotFoundError(f"Source CSV not found: {csv_path}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Section 10 Case Law Database"

    # Read CSV data
    rows = []
    with open(csv_path, encoding="utf-8") as f:
        reader = csv.reader(f)
        for r in reader:
            rows.append(r)

    if not rows:
        print("No rows found in CSV!")
        return

    # Header styling: Navy fill, white bold text
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Data styling
    data_font = Font(name="Calibri", size=10)
    data_alignment_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)
    data_alignment_standard = Alignment(horizontal="left", vertical="top")

    # Thin light gray borders
    thin_border = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB")
    )

    # Write rows
    for row_idx, row in enumerate(rows, 1):
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            if row_idx == 1:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
            else:
                cell.font = data_font
                # Wrap text for long text columns
                if col_idx in (2, 6, 7, 8, 10, 11, 12, 13, 14, 15, 16, 17):
                    cell.alignment = data_alignment_wrap
                else:
                    cell.alignment = data_alignment_standard

    # Freeze header row
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 28

    # Set custom column widths
    column_widths = {
        1: 22,  # Court
        2: 36,  # Case Title
        3: 20,  # Case Number
        4: 15,  # Judgment Date
        5: 25,  # Judges
        6: 30,  # Acts Cited
        7: 25,  # Sections Cited
        8: 25,  # Legal Topics
        9: 25,  # Outcome
        10: 42, # Fact/ Petitioner
        11: 42, # Fact/ Respondent
        12: 38, # Issues
        13: 38, # Petitioner Argument
        14: 38, # Respondent Argument
        15: 42, # Factors to be Considered
        16: 38, # Decision
        17: 35, # Citation link
        18: 20, # Ground Primary
        19: 24, # All Grounds
        20: 14, # Judgment Year
        21: 15, # Relief Status
        22: 18, # Marriage Dissolved
        23: 22, # Legal Era
        24: 22, # Substantive Adjudication
        25: 35, # Equivalent Citations
        26: 16, # Doc ID
        27: 35, # URL
        28: 16  # Chunks
    }

    for col_idx, width in column_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    print(f"Saved Excel file to: {output_path} ({os.path.getsize(output_path)} bytes)")

    # Also save to Desktop
    desktop_xlsx = "/Users/ananyasridhar/Desktop/Section10_Cases_Structured_Database.xlsx"
    try:
        wb.save(desktop_xlsx)
        print(f"Saved Desktop copy to: {desktop_xlsx}")
    except Exception as e:
        print(f"Desktop copy notice: {e}")

if __name__ == "__main__":
    build_excel()
